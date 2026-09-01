import openpyxl, os

folder = r"c:\Users\Administrator\Documents\GitHub\INtranet\uploads"
files = ["ACCESS SOH.xlsx", "KOG SOH.xlsx", "PROMED SOH.xlsx", "SEB SOH.xlsx"]

for f in files:
    path = os.path.join(folder, f)
    print("=" * 80)
    print("FILE:", f)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print("  ERROR:", e)
        continue
    print("  Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"  --- Sheet '{sheetname}' dims={ws.max_row}x{ws.max_column} ---")
        rows = list(ws.iter_rows(min_row=1, max_row=6, values_to_rows=None)) if False else None
        # print first 5 rows
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
            print(f"    row{i+1}:", [ (str(c)[:22] if c is not None else "") for c in row ])
    wb.close()
